# -*- coding: utf-8 -*-
"""Export Figures 1, 3, 4 from the revised deck: one single-slide deck per figure with the slide cropped to the figure
(so content outside the original slide is kept), PowerPoint COM export to PNG (about 600 dpi at 14.65 cm) and PDF,
blank margins cropped, dpi tag set for a 14.65 cm printed width."""
import os, sys, copy, subprocess, shutil, importlib
from pptx import Presentation
from pptx.util import Emu
from PIL import Image
import numpy as np

import schem_paths as _SP                               # repository paths
DECK = str(_SP.DECK)
OUTF = str(_SP.FIG_OUT)
EXP = str(_SP.EXPORT)
os.makedirs(EXP, exist_ok=True)
FIGS = {'1': 1, '3': 2, '4': 3}   # figure -> slide number of the revised deck (build_all.py keeps only these three slides)
PRINT_W_CM = 14.65
TARGET_PX = 3460
CM = 360000

which = sys.argv[1:] or ['1', '3', '4']


def single_slide_deck(fig, slide_no):
    mod = importlib.import_module('build_fig%s' % fig.split('_')[0])
    x0, y0, x1, y1 = mod.CROP
    prs = Presentation(DECK)
    sldIdLst = prs.slides._sldIdLst
    assert len(prs.slides) == 3, len(prs.slides)        # build_all.py: Figs. 1, 3, 4 only
    for i, sldId in reversed(list(enumerate(list(sldIdLst)))):
        if i != slide_no - 1:
            prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
    s = prs.slides[0]
    dx, dy = int(round(x0 * CM)), int(round(y0 * CM))
    for sh in s.shapes:
        sh.left = sh.left - dx
        sh.top = sh.top - dy
    prs.slide_width = Emu(int(round((x1 - x0) * CM)))
    prs.slide_height = Emu(int(round((y1 - y0) * CM)))
    p = os.path.join(EXP, 'Figure_%s.pptx' % fig)
    prs.save(p)
    return p, (x1 - x0), (y1 - y0)


PS = r'''
$ErrorActionPreference = "Stop"
$pp = New-Object -ComObject PowerPoint.Application
try {
  $pres = $pp.Presentations.Open("%(pptx)s", $true, $false, $false)
  $pres.Slides.Item(1).Export("%(png)s", "PNG", %(w)d, %(h)d)
  $pres.SaveAs("%(pdf)s", 32)
  $pres.Close()
} finally {
  if ($pp.Presentations.Count -eq 0) { $pp.Quit() | Out-Null }   # never close a PowerPoint session someone else uses
  [System.Runtime.Interopservices.Marshal]::ReleaseComObject($pp) | Out-Null
}
'''


def trim(png_in, png_out, pad_px=14):
    im = Image.open(png_in).convert('RGB')
    a = np.asarray(im).astype(int)
    nonwhite = (a < 245).any(axis=2)
    ys, xs = np.where(nonwhite)
    x0, x1 = max(xs.min() - pad_px, 0), min(xs.max() + pad_px + 1, im.width)
    y0, y1 = max(ys.min() - pad_px, 0), min(ys.max() + pad_px + 1, im.height)
    im = im.crop((x0, y0, x1, y1))
    dpi = im.width / (PRINT_W_CM / 2.54)
    tmp = os.path.join(EXP, os.path.basename(png_out))           # write next to the export decks, then copy (retry if the
    im.save(tmp, dpi=(dpi, dpi))                                  # target is briefly held open by another process)
    import time
    for k in range(30):
        try:
            shutil.copyfile(tmp, png_out)
            break
        except OSError:
            time.sleep(2)
    else:
        raise OSError('could not write ' + png_out)
    return im.size, dpi


for fig in which:
    slide_no = FIGS[fig]
    pptx, wcm, hcm = single_slide_deck(fig, slide_no)
    wpx = int(round(TARGET_PX * 1.0))
    hpx = int(round(wpx * hcm / wcm))
    raw = os.path.join(EXP, 'Figure_%s_raw.png' % fig)
    pdf = os.path.join(OUTF, 'Figure_%s.pdf' % fig)
    ps1 = os.path.join(EXP, 'export_%s.ps1' % fig)
    open(ps1, 'w', encoding='utf-8').write(PS % dict(pptx=pptx.replace('/', '\\'), png=raw.replace('/', '\\'),
                                                       pdf=pdf.replace('/', '\\'), w=wpx, h=hpx))
    r = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', ps1], capture_output=True, text=True)
    if r.returncode:
        print(r.stdout, r.stderr)
        raise SystemExit('export failed for figure %s' % fig)
    size, dpi = trim(raw, os.path.join(OUTF, 'Figure_%s.png' % fig))
    print('Figure_%s: slide %.2f x %.2f cm, png %s, %.0f dpi at %.2f cm; print scale %.3f' % (fig, wcm, hcm, size, dpi, PRINT_W_CM,
                                                                                           PRINT_W_CM / wcm))
