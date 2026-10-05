# -*- coding: utf-8 -*-
"""Rebuild out/figures/Ppt_edit_figures_REVISED.pptx from a fresh copy of the original deck and apply the edits of
Figs. 1, 3 and 4 (original slides 1, 4 and 5). The original deck in figures_src/ is only read.
Author decision of 2 Oct 2026 (the mine operator does not permit photographs): Fig. 1 is built without photographs
(build_fig1.edit), and the revised deck keeps ONLY the three figure slides, in the order Fig. 1, Fig. 3, Fig. 4
(slides 1, 2, 3; export_figs.FIGS). The unused slides of the original deck (old figures, some on the bench photographs)
are removed, and every image relationship that no shape references any more is dropped, so that the saved deck holds no
photograph (check: python check_media.py)."""
import shutil, sys, importlib
from pptx import Presentation

import schem_paths as _SP                               # repository paths
SRC = str(_SP.original_deck())                          # the authors' original deck (not released)
DST = str(_SP.DECK)
KEEP = (1, 4, 5)                                   # original slide numbers of Figs. 1, 3, 4
RT_IMAGE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
RT_MEDIA = ('/relationships/image', '/relationships/media', '/relationships/video', '/relationships/audio', '/relationships/oleObject',
            '/relationships/package')


def drop_unused_rels(slide):
    """drop picture/media relationships of the slide that its XML no longer references (deleted pictures)"""
    xml = slide._element.xml
    gone = []
    for rId, rel in list(slide.part.rels.items()):
        if rel.reltype.endswith(RT_MEDIA) and ('"%s"' % rId) not in xml:
            slide.part.drop_rel(rId); gone.append(rId)
    return gone


which = sys.argv[1:] or ['1', '3', '4']
assert set(which) == {'1', '3', '4'}, 'the deck is always rebuilt with all three figures (unused slides are removed)'
shutil.copyfile(SRC, DST)
prs = Presentation(DST)
for n in which:
    mod = importlib.import_module('build_fig%s' % n)
    mod.edit(prs)
    print('edited figure', n)
lst = prs.slides._sldIdLst
for i, sldId in reversed(list(enumerate(list(lst)))):
    if i + 1 not in KEEP:
        prs.part.drop_rel(sldId.rId)
        lst.remove(sldId)
for k, sl in enumerate(prs.slides, 1):
    print('slide %d: unreferenced media relationships dropped: %s' % (k, drop_unused_rels(sl) or 'none'))
assert len(prs.slides) == 3
prs.save(DST)
print('saved', DST, '(%d slides: Figs. 1, 3, 4)' % len(prs.slides))
