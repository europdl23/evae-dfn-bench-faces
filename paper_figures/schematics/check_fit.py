"""Text-fit and font audit of Figs. 1, 3, 4 (slides 1, 2, 3 of the revised deck; no-photo Fig. 1 since 2 Oct 2026).
2 Oct 2026: smallest printed size counts sub- and superscripts at their drawn size (pp.fit_report); no-photo slide added."""
import sys, importlib
from pptx import Presentation
import pp
import schem_paths as _SP                               # repository paths
DECK = str(_SP.DECK)
prs = Presentation(DECK)
assert len(prs.slides) == 3, len(prs.slides)
FIGS = (('1', 1), ('3', 2), ('4', 3))
which = sys.argv[1:] or [f for f, _ in FIGS]
for fig, sl in FIGS:
    if fig not in which:
        continue
    crop = importlib.import_module('build_fig%s' % fig.split('_')[0]).CROP
    scale = 14.65 / (crop[2] - crop[0])
    issues, smallest = pp.fit_report(prs.slides[sl - 1], scale=scale, crop=crop)
    print('Figure %s (slide %d): print scale %.3f, smallest printed text %.2f pt, %d issues' % (fig, sl, scale, smallest, len(issues)))
    for i in issues:
        print('   ', i)

# font audit: every run must be Times New Roman (explicitly)
for fig, sl in FIGS:
    if fig not in which:
        continue
    bad = []
    for sh in prs.slides[sl - 1].shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() and r.font.name != 'Times New Roman':
                        bad.append((sh.name, r.text[:20], r.font.name))
    print('Figure %s: runs not in Times New Roman: %d %s' % (fig, len(bad), bad[:5]))
