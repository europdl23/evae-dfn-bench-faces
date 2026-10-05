# -*- coding: utf-8 -*-
"""Small python-pptx helpers for the methodology schematics (positions in cm, fonts Times New Roman)."""
import copy
from lxml import etree
from pptx.util import Emu, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from PIL import ImageFont

CM = 360000
FONT = 'Times New Roman'
PAL = dict(green_hdr='C2F1C8', green_box='47D45A', magenta='D86ECC', orange='FFC000', blue='4E95D9',
           navy='042433', grey='F2F2F2', grey2='D9D9D9', black='000000', white='FFFFFF', lightblue='C1E5F5',
           teal='156082', evae='009E73', adfne='CC79A7', kde='666666', set1='0072B2', set2='D55E00')


def E(cm):
    return Emu(int(round(cm * CM)))


def cm(emu):
    return emu / CM


# ---------------------------------------------------------------- text measurement (Times New Roman)
_FONTS = {}


def _font(bold=False, italic=False):
    key = (bold, italic)
    if key not in _FONTS:
        f = {(False, False): 'times.ttf', (True, False): 'timesbd.ttf', (False, True): 'timesi.ttf', (True, True): 'timesbi.ttf'}[key]
        _FONTS[key] = ImageFont.truetype('C:/Windows/Fonts/' + f, 200)
    return _FONTS[key]


def text_w(s, pt, bold=False, italic=False):
    """width in cm of string s at pt points"""
    w = _font(bold, italic).getlength(s)          # at 200 px em
    return w / 200.0 * pt / 72.0 * 2.54


def line_h(pt, spacing=1.0):
    return pt / 72.0 * 2.54 * 1.2 * spacing


# ---------------------------------------------------------------- shapes
def _set_line(shape_line, color=None, lw=None, dash=None):
    if color is None:
        shape_line.fill.background()
        return
    shape_line.color.rgb = RGBColor.from_string(color)
    if lw is not None:
        shape_line.width = Pt(lw)
    if dash:
        ln = shape_line._get_or_add_ln()
        for d in ln.findall(qn('a:prstDash')):
            ln.remove(d)
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', dash)
        # prstDash must come after fill in schema order; move it right after solidFill
        sf = ln.find(qn('a:solidFill'))
        if sf is not None:
            sf.addnext(pd)


def _clear_style_font(shape):
    pass


def box(slide, x, y, w, h, fill=None, line=None, lw=1.0, dash=None, kind='rect', adj=None, name=None):
    k = {'rect': MSO_SHAPE.RECTANGLE, 'round': MSO_SHAPE.ROUNDED_RECTANGLE, 'oval': MSO_SHAPE.OVAL,
         'rbracket': MSO_SHAPE.RIGHT_BRACKET, 'lbracket': MSO_SHAPE.LEFT_BRACKET, 'rarrow': MSO_SHAPE.RIGHT_ARROW,
         'darrow': MSO_SHAPE.DOWN_ARROW, 'arc': MSO_SHAPE.ARC}[kind]
    s = slide.shapes.add_shape(k, E(x), E(y), E(w), E(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    _set_line(s.line, line, lw, dash)
    if adj is not None:
        for i, v in enumerate(adj if isinstance(adj, (list, tuple)) else [adj]):
            s.adjustments[i] = v
    # no shadow
    sp = s._element.spPr
    if sp.find(qn('a:effectLst')) is None:
        etree.SubElement(sp, qn('a:effectLst'))
    if name:
        s.name = name
    return s


def _apply_run(r, size, bold=False, italic=False, color='000000', sub=False, sup=False, font=FONT):
    f = r.font
    f.name = font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    for tag in ('a:cs', 'a:ea'):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set('typeface', font)
    if sub:                                   # True: PowerPoint's default offset; an int: offset in 1/1000 % of the run size
        rPr.set('baseline', '-25000' if sub is True else str(-abs(int(sub))))
    if sup:
        rPr.set('baseline', '30000')


def set_paragraphs(tf, paras, size, align='ctr', color='000000', bold=False, italic=False, spacing=None):
    """paras: list of paragraphs; each paragraph a string or a list of (text, opts dict) runs"""
    al = {'ctr': PP_ALIGN.CENTER, 'l': PP_ALIGN.LEFT, 'r': PP_ALIGN.RIGHT}[align]
    # clear
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p0 = tf.paragraphs[0]
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for i, para in enumerate(paras):
        p = p0 if i == 0 else tf.add_paragraph()
        p.alignment = al
        if spacing:
            p.line_spacing = spacing
        runs = [(para, {})] if isinstance(para, str) else para
        for txt, o in runs:
            r = p.add_run()
            r.text = txt
            _apply_run(r, o.get('size', size), o.get('bold', bold), o.get('italic', italic), o.get('color', color),
                       o.get('sub', False), o.get('sup', False))
        # end paragraph props so empty lines keep size
        epr = p._p.get_or_add_endParaRPr()
        epr.set('sz', str(int(size * 100)))


def text(slide, x, y, w, h, paras, size, align='ctr', anchor='ctr', color='000000', bold=False, italic=False,
         fill=None, line=None, lw=1.0, dash=None, kind='rect', margin=0.05, spacing=None, name=None, wrap=True):
    if fill is None and line is None and kind == 'rect':
        s = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    else:
        s = box(slide, x, y, w, h, fill=fill, line=line, lw=lw, dash=dash, kind=kind)
    tf = s.text_frame
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = E(margin)
    tf.margin_top = tf.margin_bottom = E(margin * 0.5)
    tf.vertical_anchor = {'ctr': MSO_ANCHOR.MIDDLE, 't': MSO_ANCHOR.TOP, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = paras.split('\n')
    set_paragraphs(tf, paras, size, align, color, bold, italic, spacing)
    if name:
        s.name = name
    return s


def arrow(slide, x1, y1, x2, y2, color='000000', lw=1.25, head='triangle', tail=None, dash=None, hw='med', hl='med'):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = RGBColor.from_string(color)
    c.line.width = Pt(lw)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn('a:prstDash'))
        pd.set('val', dash)
    if tail:
        t = etree.SubElement(ln, qn('a:headEnd'))
        t.set('type', tail); t.set('w', hw); t.set('len', hl)
    if head:
        t = etree.SubElement(ln, qn('a:tailEnd'))
        t.set('type', head); t.set('w', hw); t.set('len', hl)
    return c


def line(slide, x1, y1, x2, y2, color='000000', lw=1.25, dash=None):
    return arrow(slide, x1, y1, x2, y2, color, lw, head=None, dash=dash)


def polyline_arrow(slide, pts, color='000000', lw=1.25, head_last=True):
    out = []
    for i in range(len(pts) - 1):
        (a, b), (c, d) = pts[i], pts[i + 1]
        out.append(arrow(slide, a, b, c, d, color, lw, head='triangle' if (head_last and i == len(pts) - 2) else None))
    return out


def picture(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, E(x), E(y), E(w) if w else None, E(h) if h else None)


def delete(shape):
    el = shape._element
    el.getparent().remove(el)


def by_id(slide, sid):
    for s in slide.shapes:
        if s.shape_id == sid:
            return s
    raise KeyError(sid)


def move(shape, x=None, y=None, w=None, h=None, dx=0.0, dy=0.0):
    if x is not None:
        shape.left = E(x)
    if y is not None:
        shape.top = E(y)
    if w is not None:
        shape.width = E(w)
    if h is not None:
        shape.height = E(h)
    if dx:
        shape.left = shape.left + E(dx)
    if dy:
        shape.top = shape.top + E(dy)


def set_font_size(shape, size):
    if not shape.has_text_frame:
        return
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = Pt(size)
        epr = p._p.find(qn('a:endParaRPr'))
        if epr is not None:
            epr.set('sz', str(int(size * 100)))


def bbox(shapes):
    xs, ys, xe, ye = [], [], [], []
    for s in shapes:
        xs.append(cm(s.left)); ys.append(cm(s.top)); xe.append(cm(s.left + s.width)); ye.append(cm(s.top + s.height))
    return min(xs), min(ys), max(xe), max(ye)


# ---------------------------------------------------------------- text-fit check (greedy wrap with Times metrics)
SUB_SCALE = 2 / 3     # PowerPoint draws a baseline-shifted (sub- or superscript) run at about 2/3 of its size (5.3 pt of
                      # 8.0 pt measured in the exported PDF of Fig. 3 before the fix of 2 Oct 2026)


def _run_eff(r):
    """printed glyph size of a run on the slide (pt), allowing for the sub/superscript reduction"""
    size = r.font.size.pt if r.font.size is not None else 12
    rPr = r._r.find(qn('a:rPr'))
    if rPr is not None and rPr.get('baseline') not in (None, '0'):
        return size * SUB_SCALE, True
    return size, False


def fit_report(slide, min_pt=None, scale=1.0, crop=None):
    """flag text that wraps beyond its box height, or words wider than the box; report the smallest printed size
    (2 Oct 2026: sub- and superscript runs count at their drawn size; an exact line spacing in points is honoured)"""
    out, smallest = [], 99
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        if crop is not None:
            x, y = cm(sh.left), cm(sh.top)
            if not (crop[0] - 0.5 <= x <= crop[2] and crop[1] - 0.5 <= y <= crop[3]):
                continue
        tf = sh.text_frame
        w = cm(sh.width) - cm(tf.margin_left or 0) - cm(tf.margin_right or 0)
        h = cm(sh.height) - cm(tf.margin_top or 0) - cm(tf.margin_bottom or 0)
        tot = 0.0
        for p in tf.paragraphs:
            effs = [_run_eff(r) for r in p.runs if r.text.strip()] or [(12, False)]
            smallest = min(smallest, min(e for e, _ in effs) * scale)
            body = [e for e, sub in effs if not sub] or [max(e for e, _ in effs)]
            pt = max(body)
            words = p.text.split(' ')
            nlines, cur = 1, ''
            for wd in words:
                trial = (cur + ' ' + wd).strip()
                if text_w(trial, pt) <= w or not cur:
                    if text_w(wd, pt) > w + 0.02:
                        out.append('WORD TOO WIDE [%s] %r in %.2f cm' % (sh.name, wd, w))
                    cur = trial
                else:
                    nlines += 1; cur = wd
            ls = p.line_spacing
            lh = (ls.pt / 72.0 * 2.54) if (ls is not None and hasattr(ls, 'pt')) else line_h(pt)
            tot += nlines * lh
        if tot > h + 0.08:
            out.append('OVERFLOW [%s] %r needs %.2f cm, box %.2f cm' % (sh.name, tf.text[:40], tot, h))
    return out, smallest


# ---------------------------------------------------------------- duplicate a slide (2 Oct 2026, Figure_1_no_photo)
def duplicate_slide(prs, index):
    """append a copy of slide `index` (0-based) at the end of the deck: shapes deep-copied, picture and other internal
    relationships re-pointed to the same parts (the source slide is not changed); returns the new slide"""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    src = prs.slides[index]
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        delete(sh)
    rmap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rmap[rId] = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rmap[rId] = new.part.relate_to(rel.target_part, rel.reltype)
    bg = src._element.cSld.find(qn('p:bg'))
    if bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(bg))
    tree = new.shapes._spTree
    for el in src.shapes._spTree:
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            continue
        e2 = copy.deepcopy(el)
        for node in e2.iter():
            for k, v in list(node.attrib.items()):
                if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}') and v in rmap:
                    node.set(k, rmap[v])
        tree.append(e2)
    return new
