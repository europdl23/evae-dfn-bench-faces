# -*- coding: utf-8 -*-
"""Small Excel helpers (openpyxl): a workbook, and a sheet with a bold wrapped header, number formats and widths."""
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


def book():
    wb = Workbook(); wb.remove(wb.active); return wb


def sheet(wb, title, header, rows, fmts=None, widths=None, freeze='A2'):
    ws = wb.create_sheet(title); ws.append(header)
    for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical='top')
    for i, r in enumerate(rows):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in r])
        if fmts:
            for c in ws[ws.max_row][2:]: c.number_format = fmts[i]
    for j, w in enumerate(widths or [], start=1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = freeze
    return ws
