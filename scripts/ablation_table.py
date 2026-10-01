# -*- coding: utf-8 -*-
"""The ablation table (docs/ABLATION.md): held-out and the four arms of the 2 x 2 (single / dual latent, without / with
the geology-informed terms) on the held-out panels of all three tests, with the same parameters as the comparison table.

  tables/ablation_table.csv        parameters per scheme for held-out and the four arms (full precision)
  tables/ablation_properties.csv   per test box: held-out value and the median over the 24 networks of each arm
  tables/ablation_tests_*.csv      the four edges of the 2 x 2, tested per score and scheme
  tables/ablation_tables.xlsx      all of it in one Excel file
  figures/Table_ablation.png/.pdf  the table as an image (16 cm wide; bold = closest to held-out in that scheme)

Tests: every network is scored against its held-out box (the 38 error scores of scripts/statistical_tests.py); box
value = median over 24 networks; for each edge, a paired two-sided Wilcoxon over the test boxes of a scheme, Holm over
the 4 edges within each score and scheme, alpha 0.05. A component matters for a score when its edge is significant in
the same direction in at least 2 of the 3 schemes.
Undefined scores are n.d. (left out of panel medians and tests), as in the paper; --penalty counts them as the worst
value instead (first version of the study; outputs end in _penalty).
Usage:  python scripts/ablation_table.py [--check] [--penalty]   (--check: compare the table and the test counts with tables/reference/)"""
import sys
from pathlib import Path
from multiprocessing import Pool
import numpy as np, pandas as pd
from scipy.stats import wilcoxon

HERE = Path(__file__).resolve().parent
sys.path[:0] = [str(HERE), str(HERE.parent / 'src' / 'study')]
import study as C              # noqa: E402
import setbg as SB             # noqa: E402
import repo_paths as RP        # noqa: E402
import ablation as AB          # noqa: E402
import make_figures as MF      # noqa: E402  props, canon, table_value, TABP (same definitions as the comparison table)
import statistical_tests as ST  # noqa: E402  score, GROUPS, NAMES
import export_excel as EX      # noqa: E402  ROWS, DEC, score_name, _freeze

ARMS = AB.ORDER
COLS = ['Held-out'] + ARMS
EDGES = [('Geology-informed terms, single latent', 'single_guide', 'single_noguide'), ('Geology-informed terms, dual latent', 'EVAE', 'dual_noguide'),
         ('Dual vs single latent, with terms', 'EVAE', 'single_guide'), ('Dual vs single latent, without terms', 'dual_noguide', 'single_noguide')]


def fold_job(job):
    s, fi = job
    lib = C.load_library(); sp = C.split(s, fi, lib); model = C.cluster_model(s, fi); rows, traces, scores = [], [], []
    for b in sp['test']:
        real = sp['panels'][b]; ym = real['y_max']
        nets = {'Held-out': [real['lines']]}
        for arm in ARMS:
            nets[arm] = [C.load_net(AB.net_path(arm, s, fi, b, sd, d)) for sd in C.SEEDS for d in range(C.NDRAW)]
        for m, Ls in nets.items():
            P = [MF.props(L, ym, model) for L in Ls]
            rows.append(dict(scheme=s, fold=fi, box=b, method=m, **(P[0] if m == 'Held-out' else pd.DataFrame(P).median(numeric_only=True).to_dict())))
            for k, L in enumerate(Ls):
                L = np.asarray(L, float).reshape(-1, 4)
                if m != 'Held-out':
                    scores.append(dict(scheme=s, fold=fi, box=b, method=m, net=k, **ST.score(L, real['lines'], ym, model)))
                if not len(L): continue
                _, ln, th = C.common.geom(L)
                traces.append(pd.DataFrame(dict(scheme=s, method=m, box=b, length=ln * C.M, theta=np.degrees(th),
                                                oc=MF.canon(SB.classify(th, model), model), weight=1.0 / len(Ls))))
    return rows, pd.concat(traces), scores


def edge_tests(BM):
    out = []
    for s in C.SCHEMES:
        sub = BM[BM.scheme == s]
        for grp, names in ST.GROUPS:
            for n in names:
                piv = sub.pivot_table(index='box', columns='method', values=n)
                res = []
                for name, a, b in EDGES:
                    q = piv[[a, b]].replace([np.inf], 1e9).dropna(); d = (q[a] - q[b]).values
                    p = 1.0 if len(d) < 3 or np.allclose(d, 0) else float(wilcoxon(d, zero_method='zsplit').pvalue)
                    res.append([name, a, b, p, float(np.median(d)) if len(d) else np.nan, len(q)])
                order = sorted(range(len(res)), key=lambda i: res[i][3]); run = 0.0
                for r_, i in enumerate(order):
                    run = max(run, min(1.0, res[i][3] * (len(res) - r_))); res[i].append(run)
                for name, a, b, p, md, nb, ph in res:
                    out.append(dict(scheme=s, group=grp, score=n, edge=name, first=a, second=b, n_boxes=nb, median_diff=md, p_holm=ph,
                                    result='no difference' if ph >= 0.05 else ('better' if md < 0 else 'worse')))
    return pd.DataFrame(out)


def table_image(TBL, FS, plt):
    """16 cm wide; numbers 8 pt in one sub-column per scheme; bold = closest to held-out among the four arms"""
    short = ['Traces per 100 m²', 'Median length (m)', '90th pct. length (m)', 'Intensity P21 (m/m²)', 'Intersections per 100 m²',
             'Connections per trace', 'Spacing, nearest (m)', 'C1 (~125°) share (%)', 'C2 (~56°) share (%)', 'C3 (~27°) share (%)',
             'C4 (~89°) share (%)', 'C1 direction (°)', 'C2 direction (°)', 'C1 spread (°)', 'C2 spread (°)']
    rh, top, nh, n = 0.5, 0.3, 3, len(MF.TABP)
    Hcm = top + (nh + n) * rh + 1.5
    fig = plt.figure(figsize=(FS.FULL_W, Hcm * FS.CM)); ax = fig.add_axes([0, 0, 1, 1]); ax.axis('off'); ax.set_xlim(0, 16); ax.set_ylim(Hcm, 0)
    x0 = 3.75; gw = (15.9 - x0) / 5                                         # first group's left edge, group width (cm)
    sub = lambda g, k: x0 + g * gw + (k + 0.5) * gw / 3                    # centre of scheme k in group g
    yA, yB, yC = top + 0.5 * rh, top + 1.5 * rh, top + 2.5 * rh
    ax.text(0.2, yB, 'Parameter', fontsize=12, fontweight='bold', va='center')
    ax.text(x0 + 0.5 * gw, yB, 'Held-out', fontsize=10, fontweight='bold', va='center', ha='center')
    for g0, lab in ((1, 'Single latent'), (3, 'Dual latent')):
        ax.text(x0 + (g0 + 1) * gw, yA, lab, fontsize=12, fontweight='bold', va='center', ha='center')
        ax.plot([x0 + g0 * gw + 0.12, x0 + (g0 + 2) * gw - 0.12], [top + rh, top + rh], color='k', lw=0.5)
    for g, lab in enumerate(['', 'No terms', 'With terms', 'No terms', 'With terms']):
        if lab: ax.text(x0 + (g + 0.5) * gw, yB, lab, fontsize=10, va='center', ha='center', color='#007A5A' if g == 4 else 'k')
        for k, s in enumerate(C.SCHEMES):
            ax.text(sub(g, k), yC, s, fontsize=8, va='center', ha='center')
    ax.plot([0.1, 15.9], [top, top], color='k', lw=0.9); ax.plot([0.1, 15.9], [top + nh * rh, top + nh * rh], color='k', lw=0.6)
    for i, ((key, lab, fmt), lab2) in enumerate(zip(MF.TABP, short)):
        y = top + (nh + i + 0.5) * rh
        ax.text(0.2, y, lab2, fontsize=10, va='center')
        r = TBL.iloc[i]
        for k, s in enumerate(C.SCHEMES):
            ref = r['%s Held-out' % s]
            err = {}
            for arm in ARMS:
                v = r['%s %s' % (s, arm)]
                e = abs((v - ref + 90) % 180 - 90) if key.startswith('dir') else abs(v - ref)
                err[arm] = float(fmt % e)
            best = min(err.values())
            ax.text(sub(0, k), y, fmt % ref, fontsize=8, va='center', ha='center')
            for g, arm in enumerate(ARMS, start=1):
                ax.text(sub(g, k), y, fmt % r['%s %s' % (s, arm)], fontsize=8, va='center', ha='center',
                        fontweight='bold' if err[arm] == best else 'normal', color='#007A5A' if arm == 'EVAE' else 'k')
    yb = top + (nh + n) * rh
    ax.plot([0.1, 15.9], [yb, yb], color='k', lw=0.9)
    for k_, line in enumerate(['Dual latent with terms = the full EVAE (green). Bold = closest to held-out in that test (ties: all bold).',
                               'Terms = geology-informed terms (geology, trace-direction and cluster terms). C1-C4 = trace-direction clusters. Same seeds in every column.',
                               'S1 fill a gap (20 test boxes); S2 new bench level (50); S3 new part of the wall (50).']):
        ax.text(0.2, yb + 0.3 + 0.36 * k_, line, fontsize=8, va='top')
    return FS.save(fig, str(RP.FIGURES_DIR / 'Table_ablation'))


def write_excel(TBL, PR, DD, CNT, MAT):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    lab = {'Held-out': 'Held-out', **{a: AB.ARMS[a][0] for a in ARMS}}
    wb = Workbook(); wb.remove(wb.active); bold = Font(bold=True)

    def sheet(title, header, rows, fmts=None, widths=None, freeze='A2'):
        ws = wb.create_sheet(title); ws.append(header)
        for c in ws[1]: c.font = bold; c.alignment = Alignment(wrap_text=True, vertical='top')
        for i, r in enumerate(rows):
            ws.append(list(r))
            if fmts:
                for c in ws[ws.max_row][2:]: c.number_format = fmts[i]
        for j, w in enumerate(widths or [], start=1): ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = freeze
        return ws
    fm = [f for _, _, f in EX.ROWS]
    head = ['Parameter', 'Unit'] + ['%s %s' % (lab[m], s) for m in COLS for s in C.SCHEMES]
    r1 = [[p, u] + [round(float(TBL.iloc[i]['%s %s' % (s, m)]), EX.DEC[f]) for m in COLS for s in C.SCHEMES] for i, (p, u, f) in enumerate(EX.ROWS)]
    sheet('Ablation table', head, r1, fm, [38, 18] + [12] * 15, 'C2')
    sheet('Tests (counts)', ['Comparison', 'Score group', 'Better', 'No difference', 'Worse'], CNT[['edge', 'group', 'better', 'no_difference', 'worse']].values.tolist(),
          None, [24, 26, 10, 14, 10])
    piv = DD.pivot_table(index=['group', 'score'], columns=['edge', 'scheme'], values='result', aggfunc='first', sort=False)
    rows = [[g, EX.score_name(sc)] + [piv.loc[(g, sc), (e, s)] for e, _, _ in EDGES for s in C.SCHEMES] for g, sc in piv.index]
    sheet('Tests per score', ['Score group', 'Score (error against held-out; smaller = closer)'] + ['%s %s' % (e, s) for e, _, _ in EDGES for s in C.SCHEMES],
          rows, None, [24, 44] + [14] * 12)
    sheet('Matters (2 of 3 schemes)', ['Score group', 'Score'] + [e for e, _, _ in EDGES],
          [[r['group'], EX.score_name(r['score'])] + [r[e] for e, _, _ in EDGES] for _, r in MAT.iterrows()], None, [24, 44] + [22] * 4)
    P2 = PR.replace({'method': lab})
    ws = sheet('Per test box', [EX.PROP_NAMES.get(c, c) for c in P2.columns], P2.where(P2.notna(), None).values.tolist(), None, [8, 6, 10, 30] + [12] * (len(P2.columns) - 4))
    for row in ws.iter_rows(min_row=2, min_col=5):
        for c in row: c.number_format = '0.000'
    sheet('Full precision', head, [[p, u] + [float(TBL.iloc[i]['%s %s' % (s, m)]) for m in COLS for s in C.SCHEMES] for i, (p, u, f) in enumerate(EX.ROWS)],
          None, [38, 18] + [12] * 15)
    notes = ['Ablation of the EVAE: single vs dual latent, without vs with the geology-informed terms (2 x 2). The dual latent with the terms is the full EVAE.',
             'Geology-informed terms = the released geology term (weight 0.005) + the trace-direction term (1.0) + the trace-direction-cluster share term (1.0). Without terms = all three off.',
             'Single latent = one joint encoder; every decoder head reads the whole 56-d code (Paper 1 SingleLatentVAE). Everything else is identical.',
             'Same data, folds and settings for every arm; run seeds 1337, 20260903 and 7; 8 draws each; every arm uses the EVAE generator seeds (seeds/ablation_network_seeds.csv).',
             'Held-out test boxes: S1 fill a gap (20), S2 new bench level (50), S3 new part of the wall (50).',
             'Values: per test box the median over its 24 networks, per scheme the median over boxes; median length, 90th percentile length, cluster shares, directions and spreads over all traces (each network weighted 1/24).',
             'Tests: paired two-sided Wilcoxon over held-out panels, Holm over the 4 comparisons, alpha 0.05; undefined scores (n.d.) left out. "better" = the first-named variant (with terms, or dual latent) closer to held-out.',
             'Matters = significant in the same direction in at least 2 of the 3 schemes.',
             'Made by scripts/ablation_table.py.']
    sheet('Notes', ['Notes'], [[n] for n in notes], None, [160])
    out = RP.TABLES_DIR / 'ablation_tables.xlsx'; wb.save(out); EX._freeze(out)
    return out


def main(check=False, penalty=False):
    import matplotlib
    matplotlib.use('Agg')
    import figstyle as FS
    import matplotlib.pyplot as plt
    T = RP.TABLES_DIR; sfx = '_penalty' if penalty else ''
    lib = C.load_library(); jobs = [(s, f) for s in C.SCHEMES for f in range(len(C.folds(s, lib)))]
    with Pool(min(12, len(jobs))) as pool:
        res = pool.map(fold_job, jobs)
    PR = pd.DataFrame([r for rr, _, _ in res for r in rr]); TR = pd.concat([t for _, t, _ in res], ignore_index=True)
    SC = pd.DataFrame([r for _, _, ss in res for r in ss])
    PR.to_csv(T / 'ablation_properties.csv', index=False)
    rows = []
    for key, lab, fmt in MF.TABP:
        r = dict(parameter=lab)
        for s in C.SCHEMES:
            for m in COLS:
                r['%s %s' % (s, m)] = MF.table_value(TR, PR, s, m, key)
        rows.append(r)
    TBL = pd.DataFrame(rows); TBL.to_csv(T / 'ablation_table.csv', index=False)
    if not penalty:                                    # n.d.: undefined values are left out of the panel medians and the tests
        SC = SC.replace([np.inf, -np.inf], np.nan)
    BM = SC.groupby(['scheme', 'fold', 'box', 'method'])[ST.NAMES].median().reset_index(); BM.to_csv(T / ('ablation_tests_box_medians%s.csv' % sfx), index=False, na_rep='n.d.')
    DD = edge_tests(BM); DD.to_csv(T / ('ablation_tests_decisions%s.csv' % sfx), index=False, na_rep='n.d.')
    cnt = []
    for e, _, _ in EDGES:
        for grp in [g for g, _ in ST.GROUPS] + ['All']:
            q = DD[DD.edge == e]; q = q if grp == 'All' else q[q.group == grp]; v = q.result.value_counts()
            cnt.append(dict(edge=e, group=grp, better=int(v.get('better', 0)), no_difference=int(v.get('no difference', 0)), worse=int(v.get('worse', 0))))
    CNT = pd.DataFrame(cnt); CNT.to_csv(T / ('ablation_tests_counts%s.csv' % sfx), index=False)
    mat = []
    for grp, names in ST.GROUPS:
        for n in names:
            r = dict(group=grp, score=n)
            for e, _, _ in EDGES:
                res_ = DD[(DD.score == n) & (DD.edge == e)].result.tolist()
                r[e] = 'helps' if res_.count('better') >= 2 else ('hurts' if res_.count('worse') >= 2 else 'no consistent effect')
            mat.append(r)
    MAT = pd.DataFrame(mat); MAT.to_csv(T / ('ablation_tests_matters%s.csv' % sfx), index=False)
    if not penalty:
        print('Table image', table_image(TBL, FS, plt))
        print('Excel', write_excel(TBL, PR, DD, CNT, MAT))
    pd.set_option('display.width', 250)
    print(CNT[CNT.group == 'All'].to_string(index=False))
    print(MAT.set_index('score')[[e for e, _, _ in EDGES]].apply(pd.Series.value_counts).fillna(0).astype(int))
    if check:
        ok = True
        for f, g in (('ablation_table.csv', 'ablation_table.csv'), ('ablation_tests_counts%s.csv' % sfx, 'ablation_tests_counts.csv' if penalty else 'ablation_tests_counts_nd.csv'),
                     ('ablation_properties.csv', 'ablation_properties.csv')):
            a_ = pd.read_csv(T / f); b_ = pd.read_csv(RP.TABLES_DIR / 'reference' / g); num = a_.select_dtypes('number').columns
            same = a_.shape == b_.shape and np.allclose(a_[num].values, b_[num].values, rtol=1e-9, atol=1e-12, equal_nan=True)
            print('CHECK %-28s %s' % (f, 'identical' if same else 'DIFFERENT')); ok &= same
        if not ok:
            sys.exit(1)


if __name__ == '__main__':
    main(check='--check' in sys.argv, penalty='--penalty' in sys.argv)
