# -*- coding: utf-8 -*-
"""All table numbers in one Excel file, for building the tables by hand: tables/final_tables.xlsx

Sheets
  Held-out vs EVAE   held-out and EVAE value per held-out scheme (S1, S2, S3), rounded as reported
  All methods        held-out, EVAE, ADFNE, KDE per scheme, rounded as reported
  Tests              EVAE against ADFNE and KDE: better / no difference / worse counts per score group
  Tests per score    the result of every score in every scheme, against ADFNE and against KDE
  Verdict            three-way winner per scheme and overall, per score
  Per test box       every test box: held-out value and the median over the 24 networks of each method
  Full precision     the two table sheets without rounding
  Notes              what the sheets hold
Reads tables/final_table_all.csv and final_properties.csv (scripts/make_figures.py) and tables/stats_*.csv
(scripts/statistical_tests.py). The file is written byte for byte the same on every run (fixed zip and document dates).

Usage:  python scripts/export_excel.py"""
import sys, re, zipfile
from pathlib import Path
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
import repo_paths as RP      # noqa: E402

SCHEMES = ('S1', 'S2', 'S3')
ORDER = ('Held-out', 'EVAE', 'ADFNE', 'KDE')
# rows of the tables, in the order of tables/final_table_all.csv: (parameter, unit, number format)
ROWS = [('Trace count (P20)', 'traces per 100 m²', '0.00'), ('Median trace length', 'm', '0.0'), ('Long traces: 90th percentile length', 'm', '0.0'),
        ('Trace intensity (P21)', 'm/m²', '0.000'), ('Intersections', 'per 100 m²', '0.00'), ('Connections per trace', '-', '0.00'),
        ('Spacing: nearest trace', 'm', '0.0'),
        ('Orientation cluster 1 (~125°): share', '%', '0'), ('Orientation cluster 2 (~56°): share', '%', '0'),
        ('Orientation cluster 3 (~27°): share', '%', '0'), ('Orientation cluster 4 (~89°): share', '%', '0'),
        ('Orientation cluster 1: direction', '°', '0'), ('Orientation cluster 2: direction', '°', '0'),
        ('Orientation cluster 1: spread', '°', '0.0'), ('Orientation cluster 2: spread', '°', '0.0')]
DEC = {'0.000': 3, '0.00': 2, '0.0': 1, '0': 0}


def score_name(n):
    m = re.match(r'oc(\d)_(.*)', n)
    if m:
        j, rest = m.groups()
        return {'w1_deg': 'Cluster %s orientation (W1, °)', 'centre_err_deg': 'Cluster %s direction error (°)', 'spread_err_deg': 'Cluster %s spread error (°)',
                'share_err': 'Cluster %s share error', 'len_w1_m': 'Cluster %s length (W1, m)', 'count_relerr': 'Cluster %s count (relative error)'}[rest] % j
    return {'orient_w1_deg': 'Orientation distribution (W1, °)', 'unclustered_share_err': 'Unclustered share error', 'cluster_share_tv': 'Cluster shares (total variation)',
            'len_w1_m': 'Length distribution (W1, m)', 'len_med_err_m': 'Median length error (m)', 'len_p90_err_m': '90th percentile length error (m)',
            'P20_relerr': 'Trace count P20 (relative error)', 'P21_relerr': 'Trace intensity P21 (relative error)', 'P22_relerr': 'Intersections (relative error)',
            'nn_w1_m': 'Spacing distribution (W1, m)', 'cluster_vmr_err': 'Clustering (variance-to-mean ratio error)', 'term_share_err': 'T-ends share error',
            'CL_err': 'Connections per trace error', 'node_tv': 'Node types I / Y / X (total variation)'}[n]


PROP_NAMES = {'scheme': 'Scheme', 'fold': 'Fold', 'box': 'Box', 'method': 'Method', 'n': 'Traces', 'P20': 'Traces per 100 m²', 'P21': 'P21 (m/m²)',
              'P22': 'Intersections per 100 m²', 'len_med': 'Median length (m)', 'len_p90': '90th percentile length (m)', 'nn_med': 'Spacing, nearest trace (m)',
              'CL': 'Connections per trace', 'pI': 'Share of I nodes', 'pY': 'Share of Y nodes', 'pX': 'Share of X nodes', 'unclustered': 'Unclustered share',
              'OC1': 'Cluster 1 (~125°) share', 'OC2': 'Cluster 2 (~56°) share', 'OC3': 'Cluster 3 (~27°) share', 'OC4': 'Cluster 4 (~89°) share'}


def _freeze(path, stamp='2026-10-01T00:00:00Z'):
    """same bytes on every run: fixed dates inside the zip and in the document properties"""
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            if name == 'docProps/core.xml':
                data = re.sub(rb'(<dcterms:(created|modified)[^>]*>)[^<]*', rb'\g<1>' + stamp.encode(), data)
            zi = zipfile.ZipInfo(name, date_time=(2026, 10, 1, 0, 0, 0)); zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def main():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    T = RP.TABLES_DIR
    TBL = pd.read_csv(T / 'final_table_all.csv'); assert len(TBL) == len(ROWS)
    wb = Workbook(); wb.remove(wb.active)
    bold = Font(bold=True)

    def sheet(title, header, rows, fmts=None, widths=None):
        ws = wb.create_sheet(title); ws.append(header)
        for c in ws[1]: c.font = bold; c.alignment = Alignment(wrap_text=True, vertical='top')
        for i, r in enumerate(rows):
            ws.append(list(r))
            if fmts:
                for c in ws[ws.max_row][2:]:
                    c.number_format = fmts[i]
        for j, w in enumerate(widths or [], start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = 'C2' if fmts else 'A2'
        return ws

    fm = [f for _, _, f in ROWS]
    r1 = [[p, u] + [round(float(TBL.iloc[i]['%s %s' % (s, m)]), DEC[f]) for s in SCHEMES for m in ('Held-out', 'EVAE')] for i, (p, u, f) in enumerate(ROWS)]
    sheet('Held-out vs EVAE', ['Parameter', 'Unit'] + ['%s %s' % (s, m) for s in SCHEMES for m in ('held-out', 'EVAE')], r1, fm, [38, 18] + [11] * 6)
    r2 = [[p, u] + [round(float(TBL.iloc[i]['%s %s' % (s, m)]), DEC[f]) for s in SCHEMES for m in ORDER] for i, (p, u, f) in enumerate(ROWS)]
    sheet('All methods', ['Parameter', 'Unit'] + ['%s %s' % (s, m) for s in SCHEMES for m in ORDER], r2, fm, [38, 18] + [10] * 12)
    if (T / 'stats_counts.csv').exists():
        CNT = pd.read_csv(T / 'stats_counts.csv')
        sheet('Tests', ['EVAE against', 'Score group', 'EVAE better', 'No difference', 'EVAE worse'],
              CNT[['against', 'group', 'better', 'no_difference', 'worse']].values.tolist(), None, [14, 26, 13, 14, 12])
        DD = pd.read_csv(T / 'stats_decisions.csv')
        piv = DD.pivot_table(index=['group', 'score'], columns=['against', 'scheme'], values='result', aggfunc='first', sort=False)
        rows = [[g, score_name(sc)] + [piv.loc[(g, sc), (ag, s)] for ag in ('ADFNE', 'KDE') for s in SCHEMES] for g, sc in piv.index]
        sheet('Tests per score', ['Score group', 'Score (error against held-out; smaller = closer)'] + ['vs %s %s' % (ag, s) for ag in ('ADFNE', 'KDE') for s in SCHEMES],
              rows, None, [24, 44] + [15] * 6)
        VR = pd.read_csv(T / 'stats_verdict.csv')
        sheet('Verdict', ['Score group', 'Score', 'S1 winner', 'S2 winner', 'S3 winner', 'Overall (clear winner in at least 2 schemes)'],
              [[r.group, score_name(r.score), r.S1_winner, r.S2_winner, r.S3_winner, r.verdict] for r in VR.itertuples()], None, [24, 44, 16, 16, 16, 22])
    PR = pd.read_csv(T / 'final_properties.csv')
    ws = sheet('Per test box', [PROP_NAMES[c] for c in PR.columns], PR.where(PR.notna(), None).values.tolist(), None, [8, 6, 10, 10] + [12] * (len(PR.columns) - 4))
    for row in ws.iter_rows(min_row=2, min_col=5):
        for c in row: c.number_format = '0.000'
    full = [[p, u] + [float(TBL.iloc[i]['%s %s' % (s, m)]) for s in SCHEMES for m in ORDER] for i, (p, u, f) in enumerate(ROWS)]
    sheet('Full precision', ['Parameter', 'Unit'] + ['%s %s' % (s, m) for s in SCHEMES for m in ORDER], full, None, [38, 18] + [12] * 12)
    notes = ['EVAE = learned. ADFNE and KDE = fitted distributions (standard 2-set orientation fit, fitted lengths, mean count).',
             'Held-out test boxes of three schemes: S1 fill a gap (20 test boxes), S2 new bench level (50), S3 new part of the wall (50).',
             'Generated values: per test box, the median over its 24 networks (3 run seeds x 8 draws); per scheme, the median over its test boxes.',
             'Exceptions: median length, 90th percentile length, cluster shares, directions and spreads are taken over all traces of all test boxes of the scheme, each network weighted 1/24.',
             'Orientation cluster = a tight group of trace directions (about +-20°) learned from the fitting traces of each fold; about 125, 56, 27 and 89° on the face.',
             'Sheets "Held-out vs EVAE" and "All methods" are rounded as reported; "Full precision" has the same numbers unrounded.',
             'Tests: paired two-sided Wilcoxon over test boxes, Holm over the 2 baselines, alpha 0.05. "better" = EVAE closer to held-out.',
             'Made by scripts/export_excel.py from tables/final_table_all.csv, final_properties.csv and stats_*.csv.']
    sheet('Notes', ['Notes'], [[n] for n in notes], None, [150])
    out = T / 'final_tables.xlsx'; wb.save(out); _freeze(out)
    print('written', out)


if __name__ == '__main__':
    main()
